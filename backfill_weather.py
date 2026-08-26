"""欠測している気象データ行を、後から Open-Meteo の実況値で埋め直すバックフィルスクリプト。

collect.py の毎時収集は、DNS障害等でその回の気象取得（Open-Meteo）だけ失敗することがある
（画像取得・Gemini分類は成功しているのに気象値だけ欠測、というケース）。
本スクリプトは以下の2種類の欠測を検出し、埋め直す。

1. 気象データのみ欠測している行（行自体はあるが気温列が空）
2. 行自体が丸ごと存在しない時刻（collect.pyが気象取得より前の段階で異常終了した場合。
   例: OS再起動でプロセスごと強制終了された等）。この場合は気象データのみで行を新規挿入し、
   現象コード列は画像が存在しないため空のままにする。

いずれも該当時刻について
1. まず Archive API（ERA5実況値、信頼性が高いが直近数日分は未反映のことがある）
2. 反映されていなければ Forecast API の過去分（past_days、直近分もほぼ即時反映）
の順で値を取得する。Archive/Forecast のAPI呼び出しは同一日であれば1回にまとめ、
複数時刻分の欠測があっても日付単位でキャッシュして再利用する（無駄なAPI呼び出しを避けるため）。

30日以上前の行がそれでも埋まらない場合は、毎日ログに個別出力せず件数のみ集計する
（恒久的に取得できない古い欠測でログが埋もれるのを防ぐため）。

collect.py と同じ時間帯に同時実行されると xlsx の書き込みが競合するため、
cron 等で定期実行する場合は毎時収集(04:00-19:00台の00分)を避けたタイミングにすること。
"""
import logging
import logging.handlers
import os
import sys
import time
from datetime import datetime, timedelta

import openpyxl
import requests
from dotenv import load_dotenv

from common import (
    BASE_DIR,
    COL_DATETIME,
    COL_TEMP,
    DATA_XLSX,
    JST,
    set_weather_measures,
)

ARCHIVE_URL = "https://archive-api.open-meteo.com/v1/archive"
FORECAST_URL = "https://api.open-meteo.com/v1/forecast"

HOURLY_FIELDS = [
    "temperature_2m",
    "precipitation",
    "wind_speed_10m",
    "wind_direction_10m",
    "shortwave_radiation",
    "sunshine_duration",
    "dew_point_2m",
    "relative_humidity_2m",
    "visibility",
]

# collect.pyの観測時間帯（毎時収集の対象時刻）と合わせている
OBSERVATION_START_HOUR = 4
OBSERVATION_END_HOUR = 19

REQUEST_TIMEOUT = 20
REQUEST_RETRIES = 2
REQUEST_RETRY_DELAY_SEC = 2
BETWEEN_DATE_DELAY_SEC = 0.5

# これより古い行がまだ埋まらない場合は、個別ログを出さず件数のみ集計する
GIVE_UP_AFTER_DAYS = 30

LOG_FILE = BASE_DIR / "logs" / "backfill_weather.log"
LOG_MAX_BYTES = 2 * 1024 * 1024
LOG_BACKUP_COUNT = 3

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.handlers.RotatingFileHandler(
            LOG_FILE, maxBytes=LOG_MAX_BYTES, backupCount=LOG_BACKUP_COUNT, encoding="utf-8"
        ),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger("backfill_weather")


def query_open_meteo(url, params):
    attempts = REQUEST_RETRIES + 1
    for attempt in range(1, attempts + 1):
        try:
            resp = requests.get(url, params=params, timeout=REQUEST_TIMEOUT)
            resp.raise_for_status()
            return resp.json()
        except requests.RequestException as e:
            log.warning("Open-Meteo request failed (%s, attempt %d/%d): %s", url, attempt, attempts, e)
            if attempt < attempts:
                time.sleep(REQUEST_RETRY_DELAY_SEC)
    return None


def extract_hourly(data, target_str):
    """data の hourly から target_str（"YYYY-MM-DDTHH:00"）時刻の値一式を取り出す。無ければNone。"""
    if not data:
        return None
    hourly = data.get("hourly", {})
    times = hourly.get("time", [])
    if target_str not in times:
        return None
    idx = times.index(target_str)

    def at(key):
        vals = hourly.get(key, [])
        return vals[idx] if idx < len(vals) else None

    return {key: at(key) for key in HOURLY_FIELDS}


def raw_to_weather(raw):
    solar_wm2 = raw.get("shortwave_radiation")
    sunshine_sec = raw.get("sunshine_duration")
    visibility_m = raw.get("visibility")
    return {
        "temperature": raw.get("temperature_2m"),
        "precipitation": raw.get("precipitation"),
        "wind_speed": raw.get("wind_speed_10m"),
        "wind_direction_deg": raw.get("wind_direction_10m"),
        "solar_radiation_mj": None if solar_wm2 is None else round(solar_wm2 * 0.0036, 3),
        "sunshine_hours": None if sunshine_sec is None else round(sunshine_sec / 3600, 3),
        "dew_point": raw.get("dew_point_2m"),
        "humidity": raw.get("relative_humidity_2m"),
        "visibility_km": None if visibility_m is None else round(visibility_m / 1000, 2),
    }


def fetch_weather_for_date(date_obj, lat, lon):
    """指定日の観測時間帯（4〜19時）全時刻分の気象値をまとめて取得する。
    Archive/Forecast のAPI呼び出しはそれぞれ日付につき最大1回のみ行う
    （同日に複数時刻分の欠測があっても、同じ日のデータを重複して取得しない）。
    戻り値: {hour(int): (weather_dict, source_label)}。取得できなかった時刻はキーに含めない。
    """
    date_str = date_obj.strftime("%Y-%m-%d")
    archive_params = {
        "latitude": lat,
        "longitude": lon,
        "start_date": date_str,
        "end_date": date_str,
        "hourly": ",".join(HOURLY_FIELDS),
        "wind_speed_unit": "ms",
        "timezone": "Asia/Tokyo",
    }
    archive_data = query_open_meteo(ARCHIVE_URL, archive_params)

    forecast_data = {}
    forecast_fetched = False

    def forecast():
        nonlocal forecast_data, forecast_fetched
        if not forecast_fetched:
            days_ago = (datetime.now(JST).date() - date_obj).days
            forecast_params = {
                "latitude": lat,
                "longitude": lon,
                "hourly": ",".join(HOURLY_FIELDS),
                "wind_speed_unit": "ms",
                "timezone": "Asia/Tokyo",
                "past_days": max(1, min(days_ago + 1, 92)),
                "forecast_days": 1,
            }
            forecast_data = query_open_meteo(FORECAST_URL, forecast_params) or {}
            forecast_fetched = True
        return forecast_data

    result = {}
    for hour in range(OBSERVATION_START_HOUR, OBSERVATION_END_HOUR + 1):
        target_str = f"{date_str}T{hour:02d}:00"
        raw = extract_hourly(archive_data, target_str)
        source = "archive"
        if raw is None or raw.get("temperature_2m") is None:
            raw2 = extract_hourly(forecast(), target_str)
            if raw2 is not None and raw2.get("temperature_2m") is not None:
                raw = raw2
                source = "forecast(past_days)"
        if raw is not None and raw.get("temperature_2m") is not None:
            result[hour] = (raw_to_weather(raw), source)
    return result


def find_missing_rows(ws):
    """気温列が空の行を「気象データ欠測」とみなして (行番号, 日時) のリストを返す。"""
    missing = []
    for row in range(2, ws.max_row + 1):
        dt = ws.cell(row=row, column=COL_DATETIME).value
        if dt is None:
            continue
        if ws.cell(row=row, column=COL_TEMP).value is None:
            missing.append((row, dt))
    return missing


def find_fully_missing_rows(ws):
    """観測時間帯(4〜19時)の中で、行自体が存在しない時刻を検出する。
    collect.pyが気象取得より前の段階で異常終了した場合（例: OS再起動によるプロセス強制終了）に
    発生しうる。初日については、実際に収集が始まった時刻より前は「欠測」として扱わない。
    """
    existing = set()
    for row in range(2, ws.max_row + 1):
        dt = ws.cell(row=row, column=COL_DATETIME).value
        if isinstance(dt, datetime):
            existing.add(dt)
    if not existing:
        return []

    first_dt = min(existing)
    now = datetime.now(JST)
    today = now.date()

    missing = []
    d = first_dt.date()
    while d <= today:
        start_hour = first_dt.hour if d == first_dt.date() else OBSERVATION_START_HOUR
        for hour in range(start_hour, OBSERVATION_END_HOUR + 1):
            if d == today and hour >= now.hour:
                continue  # 現在時刻分はまだ収集中/未実行の可能性があるため対象外
            dt = datetime(d.year, d.month, d.day, hour)
            if dt not in existing:
                missing.append(dt)
        d += timedelta(days=1)
    return missing


def find_insert_position(ws, dt_naive):
    """dt_naiveを時系列順に挿入すべき行番号（既存行の直後）を返す。"""
    insert_at = 2
    for row in range(2, ws.max_row + 1):
        existing_dt = ws.cell(row=row, column=COL_DATETIME).value
        if not isinstance(existing_dt, datetime):
            continue
        if existing_dt < dt_naive:
            insert_at = row + 1
        else:
            break
    return insert_at


def main():
    load_dotenv(BASE_DIR / ".env")
    dry_run = "--dry-run" in sys.argv[1:]

    lat = os.environ.get("WEATHER_LAT", "37.34861")
    lon = os.environ.get("WEATHER_LON", "139.31583")

    if not DATA_XLSX.exists():
        log.info("xlsx not found, nothing to backfill: %s", DATA_XLSX)
        return

    log.info("=== Backfill started%s ===", " (dry-run)" if dry_run else "")

    wb = openpyxl.load_workbook(DATA_XLSX)
    ws = wb.active

    today = datetime.now(JST).date()
    date_cache = {}

    def weather_for(dt_naive):
        key = dt_naive.date().isoformat()
        if key not in date_cache:
            date_cache[key] = fetch_weather_for_date(dt_naive.date(), lat, lon)
            time.sleep(BETWEEN_DATE_DELAY_SEC)
        return date_cache[key].get(dt_naive.hour)

    gave_up = 0

    # 1) 既存行のうち気象データのみ欠測している行を埋める
    missing = find_missing_rows(ws)
    filled = 0
    if missing:
        log.info("Found %d row(s) with missing weather data", len(missing))
        for row, dt in missing:
            entry = weather_for(dt)
            if entry is None:
                if (today - dt.date()).days <= GIVE_UP_AFTER_DAYS:
                    log.info("Still unavailable, skipped: row %d (%s)", row, dt)
                else:
                    gave_up += 1
                continue
            weather, source = entry
            if not dry_run:
                set_weather_measures(ws, row, weather)
            filled += 1
            log.info("Backfilled: row %d (%s) via %s%s", row, dt, source, " [dry-run]" if dry_run else "")
    else:
        log.info("No missing weather rows found (%d data row(s) checked)", ws.max_row - 1)

    if filled and not dry_run:
        wb.save(DATA_XLSX)

    # 2) 行自体が丸ごと存在しない時刻を検出し、気象データのみで行を新規挿入する
    #    （現象コード列は画像が存在しないため空のままにする）
    fully_missing = find_fully_missing_rows(ws)
    inserted = 0
    if fully_missing:
        log.warning(
            "Found %d entirely missing row(s) (collect.py likely aborted before writing): %s",
            len(fully_missing),
            ", ".join(dt.strftime("%Y-%m-%d %H:%M") for dt in fully_missing),
        )
        for dt in fully_missing:
            entry = weather_for(dt)
            if entry is None:
                if (today - dt.date()).days <= GIVE_UP_AFTER_DAYS:
                    log.info("Entirely missing row still unavailable: %s", dt)
                else:
                    gave_up += 1
                continue
            weather, source = entry
            if not dry_run:
                insert_at = find_insert_position(ws, dt)
                ws.insert_rows(insert_at)
                ws.cell(row=insert_at, column=COL_DATETIME).value = dt
                set_weather_measures(ws, insert_at, weather)
            inserted += 1
            log.warning(
                "Inserted weather-only row for entirely missing time (phenomenon codes left blank): "
                "%s via %s%s",
                dt, source, " [dry-run]" if dry_run else "",
            )
        if inserted and not dry_run:
            wb.save(DATA_XLSX)

    if gave_up:
        log.info(
            "%d row(s) older than %d days still unavailable; will keep checking quietly",
            gave_up, GIVE_UP_AFTER_DAYS,
        )

    log.info(
        "=== Backfill finished: %d weather value(s) filled, %d fully-missing row(s) inserted ===",
        filled, inserted,
    )


if __name__ == "__main__":
    try:
        main()
    except Exception:
        log.exception("backfill_weather.py terminated abnormally")
        sys.exit(1)
