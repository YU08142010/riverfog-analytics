"""代表地点の気象実測値・現象コードのSQLite保存(只見川.xlsxの後継データストア)。

以前はcollect.py(毎時収集)・backfill_weather.py(欠測の後埋め)・backfill_images.py
(現象コードの後埋め)の3スクリプトが、同じ只見川.xlsxをopenpyxlで直接読み書きしていた。
xlsxはopenpyxlのread_onlyモードでも「全行を読み直す」以外の部分読み込み手段が無いため、
行数が増えるほど1回の読み込みコストが線形に増え続け(19,000行超で1回10秒以上)、また
4者が同じファイルの排他ロック(xlsx_write_lock)を奪い合う構造だった。

時系列データという実体はここ(SQLite)に置き、只見川.xlsx自体は「人がExcelでそのまま
見たい時にその場で書き出すエクスポート形式」という位置づけに変える(export_to_xlsx参照)。
書き込みはstation_weather.pyと同じ、datetime単位のUPSERTパターンを踏襲する。
"""
import csv
import sqlite3
from datetime import datetime as _datetime

import pandas as pd

SCHEMA = """
CREATE TABLE IF NOT EXISTS weather_observations (
    datetime TEXT PRIMARY KEY,
    temperature REAL,
    precipitation REAL,
    wind_speed REAL,
    wind_direction_deg REAL,
    solar_radiation_mj REAL,
    sunshine_hours REAL,
    dew_point REAL,
    humidity REAL,
    visibility_km REAL
);
CREATE TABLE IF NOT EXISTS phenomena (
    datetime TEXT NOT NULL,
    station_name TEXT NOT NULL,
    code INTEGER,
    PRIMARY KEY (datetime, station_name)
);
CREATE INDEX IF NOT EXISTS idx_phenomena_datetime ON phenomena(datetime);
"""

WEATHER_COLUMNS = [
    "temperature", "precipitation", "wind_speed", "wind_direction_deg",
    "solar_radiation_mj", "sunshine_hours", "dew_point", "humidity", "visibility_km",
]


def connect(db_path):
    # collect.py(毎時)・backfill_weather.py・backfill_images.py(いずれも日次)・Webの
    # 4者が同じDBファイルに書き込みうる(以前のxlsx_write_lockが解消していた競合と同じ状況)。
    # busy_timeoutを設定し、他プロセスが書き込み中でも即座にエラーにせず一定時間待たせる。
    conn = sqlite3.connect(str(db_path), timeout=30.0)
    conn.execute("PRAGMA busy_timeout = 30000")
    conn.execute("PRAGMA journal_mode = WAL")
    conn.executescript(SCHEMA)
    return conn


def record_weather(conn, dt_str, weather):
    """weather: collect.py/backfill_weather.pyのweather dict(該当キーが無い/Noneなら
    その項目は空欄)。同じdatetimeは上書き(UPSERT)する。"""
    weather = weather or {}
    cols_sql = ", ".join(WEATHER_COLUMNS)
    placeholders = ", ".join(["?"] * len(WEATHER_COLUMNS))
    updates_sql = ", ".join(f"{c}=excluded.{c}" for c in WEATHER_COLUMNS)
    with conn:
        conn.execute(
            f"""INSERT INTO weather_observations (datetime, {cols_sql})
               VALUES (?, {placeholders})
               ON CONFLICT(datetime) DO UPDATE SET {updates_sql}""",
            (dt_str, *(weather.get(c) for c in WEATHER_COLUMNS)),
        )


def record_phenomena(conn, dt_str, phenom_by_name):
    """phenom_by_name: {地点名: コード}。0="/"(現象なし)、1〜10=現象コード、
    None=未分類(空欄)。collect.pyのphenom_by_nameと同じ表現をそのまま受け取る。"""
    rows = [(dt_str, name, (None if code is None else int(code))) for name, code in phenom_by_name.items()]
    if not rows:
        return
    with conn:
        conn.executemany(
            """INSERT INTO phenomena (datetime, station_name, code) VALUES (?, ?, ?)
               ON CONFLICT(datetime, station_name) DO UPDATE SET code=excluded.code""",
            rows,
        )


def weather_row_exists(conn, dt_str):
    """指定datetimeの気象観測行が既に存在するか。backfill_images.pyが「まだ一度も
    収集されていない時刻」に誤って現象コードだけ記録してしまわないためのガード用。"""
    row = conn.execute("SELECT 1 FROM weather_observations WHERE datetime = ?", (dt_str,)).fetchone()
    return row is not None


def load_weather_df(conn):
    """datetime昇順のDataFrameを返す(列名はWEATHER_COLUMNSのまま。
    fog_features.py列名への変換はfog_data.py側で行う)。"""
    df = pd.read_sql_query(
        f"SELECT datetime, {', '.join(WEATHER_COLUMNS)} FROM weather_observations ORDER BY datetime",
        conn,
    )
    df["datetime"] = pd.to_datetime(df["datetime"])
    return df


def load_phenomena_wide_df(conn, station_names):
    """地点名を列にしたワイド形式のDataFrame(datetime, 地点名1, 地点名2, ...)を返す。
    セルの値は只見川.xlsxと同じ表現(現象なし="/"文字列、1〜10はint、未分類はNone)にし、
    fog_features.encode_phenomena_cellをそのまま使う既存呼び出し元との互換性を保つ。
    """
    df = pd.read_sql_query("SELECT datetime, station_name, code FROM phenomena", conn)
    if df.empty:
        wide = pd.DataFrame({"datetime": pd.Series(dtype="datetime64[ns]")})
        for name in station_names:
            wide[name] = pd.Series(dtype="object")
        return wide
    df["datetime"] = pd.to_datetime(df["datetime"])
    df["value"] = df["code"].map(lambda c: None if pd.isna(c) else ("/" if int(c) == 0 else int(c)))
    wide = df.pivot(index="datetime", columns="station_name", values="value").reset_index()
    wide = wide.sort_values("datetime").reset_index(drop=True)
    for name in station_names:
        if name not in wide.columns:
            wide[name] = None
    return wide[["datetime", *station_names]]


def import_from_xlsx(conn, xlsx_path, cameras_csv):
    """既存の只見川.xlsxからこのSQLiteストアへ移行する(初回移行専用)。
    UPSERTなので同じ行を何度読み直しても安全(冪等)。戻り値: (処理した行数, 地点数)。
    common.pyの固定列レイアウトをそのまま前提にする(fog_data.pyの読み込みと同じ前提)。
    """
    import openpyxl

    from common import (
        COL_DATETIME, COL_DEWPOINT, COL_HUMIDITY, COL_PHENOM_START, COL_PRECIP,
        COL_SOLAR, COL_SUNSHINE, COL_TEMP, COL_VISIBILITY, COL_WIND_DIR, COL_WIND_SPEED,
        WIND_DIR_DEGREES,
    )

    with open(cameras_csv, encoding="utf-8", newline="") as f:
        station_names = [row["地点名"] for row in csv.DictReader(f)]
    n_cols = COL_PHENOM_START - 1 + len(station_names)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.active
    n_rows = 0
    for values in ws.iter_rows(min_row=1, max_col=n_cols, values_only=True):
        dt = values[COL_DATETIME - 1]
        if not isinstance(dt, _datetime):
            continue
        dt_str = dt.strftime("%Y-%m-%d %H:%M:%S")

        wind_dir_ja = values[COL_WIND_DIR - 1]
        wind_dir_deg = WIND_DIR_DEGREES.get(str(wind_dir_ja).strip()) if wind_dir_ja else None
        record_weather(conn, dt_str, {
            "temperature": values[COL_TEMP - 1],
            "precipitation": values[COL_PRECIP - 1],
            "wind_speed": values[COL_WIND_SPEED - 1],
            "wind_direction_deg": wind_dir_deg,
            "solar_radiation_mj": values[COL_SOLAR - 1],
            "sunshine_hours": values[COL_SUNSHINE - 1],
            "dew_point": values[COL_DEWPOINT - 1],
            "humidity": values[COL_HUMIDITY - 1],
            "visibility_km": values[COL_VISIBILITY - 1],
        })

        phenom_by_name = {}
        for i, name in enumerate(station_names):
            raw = values[COL_PHENOM_START - 1 + i]
            if raw is None:
                phenom_by_name[name] = None
            elif str(raw).strip() == "/":
                phenom_by_name[name] = 0
            else:
                try:
                    phenom_by_name[name] = int(float(raw))
                except (TypeError, ValueError):
                    phenom_by_name[name] = None
        record_phenomena(conn, dt_str, phenom_by_name)
        n_rows += 1
    wb.close()
    return n_rows, len(station_names)


def export_to_xlsx(conn, dest_path, cameras_csv, template_xlsx, station_names=None):
    """このSQLiteストアの全データを、只見川.xlsxと同じ列レイアウトのxlsxとして書き出す。
    QI(品質情報)・均質番号は元データに常に一定値(QUALITY_INFO_VALUE/HOMOGENEITY_NUMBER_VALUE)
    しか入っておらず、fog_data.py等のどの読み込み処理からも参照されていないため、
    ここでは「値があれば付与、無ければ空欄」という同じ規則で書き出し時に再現する。
    station_namesを渡すと、cameras.csv全体ではなくその地点名だけを対象に書き出す
    (2026-09の会津地方追加で只見川.xlsx/会津.xlsxを地域ごとに分けて出力するため。
    渡す場合、対応するtemplate_xlsxもその地点名の並びに合わせたヘッダーである必要がある)。
    """
    import shutil

    import openpyxl

    from common import COL_DATETIME, COL_PHENOM_START, set_weather_measures

    if station_names is None:
        with open(cameras_csv, encoding="utf-8", newline="") as f:
            station_names = [row["地点名"] for row in csv.DictReader(f)]

    shutil.copy(template_xlsx, dest_path)
    wb = openpyxl.load_workbook(dest_path)
    ws = wb.active

    weather_df = load_weather_df(conn)
    phenom_df = load_phenomena_wide_df(conn, station_names)
    merged = weather_df.merge(phenom_df, on="datetime", how="outer").sort_values("datetime")

    for offset, (_, row) in enumerate(merged.iterrows(), start=2):  # 1行目はヘッダ
        r = offset
        dt = row["datetime"].to_pydatetime()
        ws.cell(row=r, column=COL_DATETIME).value = dt
        weather = {
            "temperature": row["temperature"], "precipitation": row["precipitation"],
            "wind_speed": row["wind_speed"], "wind_direction_deg": row["wind_direction_deg"],
            "solar_radiation_mj": row["solar_radiation_mj"], "sunshine_hours": row["sunshine_hours"],
            "dew_point": row["dew_point"], "humidity": row["humidity"],
            "visibility_km": row["visibility_km"],
        }
        weather = {k: (None if pd.isna(v) else v) for k, v in weather.items()}
        set_weather_measures(ws, r, weather)
        for i, name in enumerate(station_names):
            value = row[name]
            cell_value = None if pd.isna(value) else value
            ws.cell(row=r, column=COL_PHENOM_START + i).value = cell_value
    wb.save(dest_path)
    return len(merged)


__all__ = [
    "connect", "record_weather", "record_phenomena", "weather_row_exists", "load_weather_df",
    "load_phenomena_wide_df", "import_from_xlsx", "export_to_xlsx",
]
