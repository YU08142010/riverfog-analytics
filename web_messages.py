"""webapp.py から分けた部分(web_messages)。webapp.py の最後で読み込まれ、ここで定義した名前は webapp にも戻される。
ルート・テンプレート・URLの名前(エンドポイント)は分ける前と同じ。"""
import io
import json
import re
import threading
import time
import zipfile
from pathlib import Path
from flask import Response, abort, jsonify, redirect, render_template, request, send_from_directory, session, url_for
from common import now_jst_naive
import users
import push_notify
from web_game import BUBBLE_STYLE_BY_ID, DM_THEME_BY_ID

from webapp import (DM_ATTACHMENT_DIR, DM_MESSAGE_MAX_LENGTH, SHARED_FILES_DIR, STATION_NAMES, _finish, _is_online, 
    _save_uploaded_file, _users_conn, _wants_json, app, log)  # noqa: F401


# --- メッセージ機能共通(DM・グループチャット両方で使う) ------------------------------
REACTION_EMOJIS = ["👍", "❤️", "😂", "😮", "😢", "🙏"]
MESSAGE_SEARCH_LIMIT = 50

# 入力中表示: 専用テーブルは持たず、_png_cache等と同じ「モジュールレベルのdict+Lock」で
# 一時的な状態として保持する(プロセス再起動で消えても実害が無いため永続化不要)。
# キーは(channel_key, user_id)。channel_keyはDMなら"dm:<conversation_id>"、
# グループチャットなら固定文字列"group"。
_typing_state = {}
_typing_state_lock = threading.Lock()
TYPING_INDICATOR_TTL_SECONDS = 6


def _mark_typing(channel_key, user_id):
    with _typing_state_lock:
        _typing_state[(channel_key, user_id)] = time.monotonic()


def _typing_user_ids(channel_key, exclude_user_id=None):
    now = time.monotonic()
    with _typing_state_lock:
        stale = [k for k, t in _typing_state.items() if now - t > TYPING_INDICATOR_TTL_SECONDS]
        for k in stale:
            del _typing_state[k]
        return [uid for (key, uid) in _typing_state if key == channel_key and uid != exclude_user_id]


_MENTION_TOKEN_RE = re.compile(r"@([^\s@]+)")


def _parse_mentions(body, name_to_id):
    """本文中の`@表示名`を、既存ユーザー名(name_to_id: {name: user_id})と照合する。
    表示名に空白は使えない前提のため、@の後ろの非空白トークンに対し、長い名前から
    順に前方一致を試す(短い名前が長い名前の一部になっているケースの誤爆を避けるため)。"""
    matched = set()
    names_by_length = sorted(name_to_id.keys(), key=len, reverse=True)
    for m in _MENTION_TOKEN_RE.finditer(body):
        token = m.group(1)
        for name in names_by_length:
            if token.startswith(name):
                matched.add(name_to_id[name])
                break
    return matched


def _reaction_summary(reactions, my_id):
    """[{"emoji":, "user_id":}, ...] を [{"emoji":, "count":, "mine": bool}, ...] に集計する
    (絵文字が最初に現れた順を保つ)。"""
    counts = {}
    order = []
    mine = set()
    for r in reactions:
        emoji = r["emoji"]
        if emoji not in counts:
            counts[emoji] = 0
            order.append(emoji)
        counts[emoji] += 1
        if r["user_id"] == my_id:
            mine.add(emoji)
    return [{"emoji": e, "count": counts[e], "mine": e in mine} for e in order]


def _reply_preview_dict(reply_msg, name_by_id):
    """返信先メッセージの要約を、送信直後のJSONレスポンスやポーリング応答にも
    その場で載せられる形にする(ページ再読み込みを待たずに表示するため)。"""
    if reply_msg is None:
        return None
    body = reply_msg["body"]
    snippet = body[:60] + ("…" if len(body) > 60 else "")
    return {
        "id": reply_msg["id"],
        "sender_name": name_by_id.get(reply_msg["sender_id"], "(退会済みユーザー)"),
        "snippet": snippet,
    }


def _attach_message_extras(messages, reactions_by_id, mentions_by_id, name_by_id, reply_by_id, my_id):
    """メッセージ辞書のリストに、リアクション集計・メンション表示名・返信先の要約を
    その場で追加する(DM・グループ両方の描画で共通に使う)。"""
    for m in messages:
        m["reactions"] = _reaction_summary(reactions_by_id.get(m["id"], []), my_id)
        mention_ids = mentions_by_id.get(m["id"], [])
        m["mentioned_names"] = [name_by_id[uid] for uid in mention_ids if uid in name_by_id]
        try:
            m["card"] = json.loads(m["card_data"]) if m.get("card_data") else None
        except (TypeError, ValueError):
            m["card"] = None
        reply = reply_by_id.get(m.get("reply_to_id"))
        if reply:
            snippet = reply["body"][:60] + ("…" if len(reply["body"]) > 60 else "")
            m["reply_preview"] = {
                "id": reply["id"],
                "sender_name": name_by_id.get(reply["sender_id"], "(退会済みユーザー)"),
                "snippet": snippet,
            }
        else:
            m["reply_preview"] = None
    return messages


@app.route("/group")
def group_home():
    """全員参加のグループチャット。参加者管理は無く、全登録ユーザーが常にメンバー。"""
    my_id = session.get("user_id")
    if not my_id:
        return redirect(url_for("login", next=request.path))
    conn = _users_conn()
    try:
        messages = users.list_group_messages(conn)
        message_ids = [m["id"] for m in messages]
        attachments_by_id = users.get_group_attachments_for_messages(conn, message_ids)
        reactions_by_id = users.get_group_reactions_for_messages(conn, message_ids)
        mentions_by_id = users.get_group_mentions_for_messages(conn, message_ids)
        all_users = users.list_users(conn)
        pinned = users.list_pinned_group_messages(conn)
        if messages:
            users.mark_group_read(conn, my_id, messages[-1]["id"])
        read_states = users.all_group_read_states(conn)
        me = users.get_user_by_id(conn, my_id)
    finally:
        conn.close()
    name_by_id = {u["id"]: u["name"] for u in all_users}
    reply_by_id = {m["id"]: m for m in messages}
    for m in messages:
        m["attachments"] = attachments_by_id.get(m["id"], [])
        m["sender_name"] = name_by_id.get(m["sender_id"], "(退会済みユーザー)")
        m["read_count"] = sum(
            1 for uid, last_read in read_states.items() if uid != m["sender_id"] and last_read >= m["id"]
        )
    _attach_message_extras(messages, reactions_by_id, mentions_by_id, name_by_id, reply_by_id, my_id)
    for p in pinned:
        p["sender_name"] = name_by_id.get(p["sender_id"], "(退会済みユーザー)")
    my_dm_theme_class = DM_THEME_BY_ID.get(me["equipped_dm_theme"], {}).get("css_class", "") if me else ""
    my_bubble_style_class = BUBBLE_STYLE_BY_ID.get(me["equipped_bubble_style"], {}).get("css_class", "") if me else ""
    return render_template(
        "group_chat.html", messages=messages, pinned=pinned, my_id=my_id,
        dm_message_max_length=DM_MESSAGE_MAX_LENGTH, reaction_emojis=REACTION_EMOJIS,
        member_names=sorted(name_by_id.values()), my_dm_theme_class=my_dm_theme_class,
        my_bubble_style_class=my_bubble_style_class,
    )


def _serializable_read_states(read_states):
    """dictのキーがintのままだとjsonify後にJS側でstringキーになるため、
    グループチャットの既読数リアルタイム更新ではキーを文字列化して明示的に渡す。"""
    return {str(uid): last_read for uid, last_read in read_states.items()}


@app.route("/group/send", methods=["POST"])
def group_send():
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    body = request.form.get("body", "").strip()[:DM_MESSAGE_MAX_LENGTH]
    reply_to_id = request.form.get("reply_to_id", type=int)
    file = request.files.get("attachment")
    has_file = bool(file and file.filename)
    if not body and not has_file:
        return _finish("group_home", "メッセージまたは添付ファイルを指定してください。", success=False)
    attachment_info = None
    if has_file:
        result, error = _save_uploaded_file(file, DM_ATTACHMENT_DIR)
        if error:
            return _finish("group_home", error, success=False)
        attachment_info = result
    created_at = now_jst_naive().strftime("%Y-%m-%d %H:%M:%S")
    conn = _users_conn()
    try:
        if reply_to_id is not None and users.get_group_message(conn, reply_to_id) is None:
            reply_to_id = None
        message_id = users.send_group_message(conn, my_id, body, created_at, reply_to_id=reply_to_id)
        attachment_id = None
        if attachment_info:
            original_name, stored_name, content_type, size = attachment_info
            attachment_id = users.add_group_attachment(
                conn, message_id, original_name, stored_name, content_type, size, created_at,
            )
        all_users = users.list_users(conn)
        name_to_id = {u["name"]: u["id"] for u in all_users if u["id"] != my_id}
        mention_ids = _parse_mentions(body, name_to_id)
        users.record_group_mentions(conn, message_id, mention_ids)
        users.mark_group_read(conn, my_id, message_id)
        reply_preview = None
        if reply_to_id is not None:
            reply_name_by_id = {u["id"]: u["name"] for u in all_users}
            reply_preview = _reply_preview_dict(users.get_group_message(conn, reply_to_id), reply_name_by_id)
        push_ids = [uid for uid in users.push_subscriber_ids(conn) if uid != my_id]
        push_settings_raw = users.all_user_settings_raw(conn) if push_ids else {}
    finally:
        conn.close()
    try:
        my_name = session.get("name", "")
        for uid in push_ids:
            mentioned = uid in mention_ids
            push_notify.notify_async(
                [uid], "mention" if mentioned else "group",
                f"{my_name}さんがグループであなたをメンションしました" if mentioned else f"グループ: {my_name}さん",
                push_notify.message_preview(push_settings_raw.get(uid, "{}"), body, bool(attachment_info)),
                url_for("group_home"), "group", ttl=86400, urgency="high" if mentioned else "normal",
            )
    except Exception as e:
        log.warning("Failed to queue group push notifications: %s", e)
    with _typing_state_lock:
        _typing_state.pop(("group", my_id), None)
    if _wants_json():
        return jsonify({
            "message": "送信しました。", "success": True,
            "id": message_id, "body": body, "created_at": created_at,
            "attachment": (
                {"id": attachment_id, "filename": attachment_info[0]} if attachment_info else None
            ),
            "reply_preview": reply_preview,
        })
    return redirect(url_for("group_home"))


@app.route("/group/poll")
def group_poll():
    """dm_pollと同じ4秒ポーリング方式。新着メッセージに加え、入力中の相手・既読状況も返す。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    after_id = request.args.get("after", 0, type=int)
    conn = _users_conn()
    try:
        messages = users.list_group_messages(conn, after_id)
        message_ids = [m["id"] for m in messages]
        attachments_by_id = users.get_group_attachments_for_messages(conn, message_ids)
        reactions_by_id = users.get_group_reactions_for_messages(conn, message_ids)
        mentions_by_id = users.get_group_mentions_for_messages(conn, message_ids)
        all_users = users.list_users(conn)
        if messages:
            users.mark_group_read(conn, my_id, messages[-1]["id"])
        read_states = users.all_group_read_states(conn)
        reply_ids = {m["reply_to_id"] for m in messages if m.get("reply_to_id")}
        reply_msgs_by_id = {rid: users.get_group_message(conn, rid) for rid in reply_ids}
    finally:
        conn.close()
    name_by_id = {u["id"]: u["name"] for u in all_users}
    for m in messages:
        m["attachments"] = attachments_by_id.get(m["id"], [])
        m["sender_name"] = name_by_id.get(m["sender_id"], "(退会済みユーザー)")
        m["reactions"] = _reaction_summary(reactions_by_id.get(m["id"], []), my_id)
        mention_ids = mentions_by_id.get(m["id"], [])
        m["mentioned_names"] = [name_by_id[uid] for uid in mention_ids if uid in name_by_id]
        try:
            m["card"] = json.loads(m["card_data"]) if m.get("card_data") else None
        except (TypeError, ValueError):
            m["card"] = None
        m["read_count"] = sum(
            1 for uid, last_read in read_states.items() if uid != m["sender_id"] and last_read >= m["id"]
        )
        m["reply_preview"] = _reply_preview_dict(reply_msgs_by_id.get(m.get("reply_to_id")), name_by_id)
    typing_names = [
        name_by_id[uid] for uid in _typing_user_ids("group", exclude_user_id=my_id) if uid in name_by_id
    ]
    return jsonify({
        "messages": messages, "typing_names": typing_names,
        "read_states": _serializable_read_states(read_states),
    })


@app.route("/group/typing", methods=["POST"])
def group_typing():
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    _mark_typing("group", my_id)
    return jsonify({"success": True})


@app.route("/group/search")
def group_search():
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    keyword = request.args.get("q", "").strip()
    if not keyword:
        return jsonify({"results": []})
    conn = _users_conn()
    try:
        matches = users.search_group_messages(conn, keyword, limit=MESSAGE_SEARCH_LIMIT)
        name_by_id = {u["id"]: u["name"] for u in users.list_users(conn)}
    finally:
        conn.close()
    for m in matches:
        m["sender_name"] = name_by_id.get(m["sender_id"], "(退会済みユーザー)")
    return jsonify({"results": matches})


@app.route("/group/attachments/<int:attachment_id>")
def group_attachment_file(attachment_id):
    if not session.get("user_id"):
        abort(403)
    conn = _users_conn()
    try:
        att = users.get_group_attachment(conn, attachment_id)
    finally:
        conn.close()
    if att is None:
        abort(404)
    return send_from_directory(
        DM_ATTACHMENT_DIR, att["stored_name"], as_attachment=True, download_name=att["filename"],
    )


@app.route("/group/<int:message_id>/delete", methods=["POST"])
def group_message_delete(message_id):
    """自分が送信したメッセージ、または管理者は誰のメッセージでも削除できる
    (荒らし・誤爆投稿を全員が見えるチャンネルから消せるようにするため)。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        msg = users.get_group_message(conn, message_id)
        if msg is None:
            abort(404)
        if msg["sender_id"] != my_id and session.get("role") != "admin":
            return jsonify({"success": False, "message": "自分が送信したメッセージのみ削除できます。"}), 403
        stored_names = users.get_group_message_attachment_stored_names(conn, message_id)
        users.delete_group_message(conn, message_id)
    finally:
        conn.close()
    for stored_name in stored_names:
        (DM_ATTACHMENT_DIR / stored_name).unlink(missing_ok=True)
    return jsonify({"success": True, "message_id": message_id})


@app.route("/group/<int:message_id>/edit", methods=["POST"])
def group_message_edit(message_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    body = request.form.get("body", "").strip()[:DM_MESSAGE_MAX_LENGTH]
    if not body:
        return jsonify({"success": False, "message": "本文を入力してください。"}), 400
    conn = _users_conn()
    try:
        msg = users.get_group_message(conn, message_id)
        if msg is None:
            abort(404)
        if msg["sender_id"] != my_id:
            return jsonify({"success": False, "message": "自分が送信したメッセージのみ編集できます。"}), 403
        edited_at = now_jst_naive().strftime("%Y-%m-%d %H:%M:%S")
        users.edit_group_message(conn, message_id, body, edited_at)
    finally:
        conn.close()
    return jsonify({"success": True, "message_id": message_id, "body": body, "edited_at": edited_at})


@app.route("/group/<int:message_id>/pin", methods=["POST"])
def group_message_pin(message_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    pinned = request.form.get("pinned") == "1"
    conn = _users_conn()
    try:
        msg = users.get_group_message(conn, message_id)
        if msg is None:
            abort(404)
        if msg["sender_id"] != my_id and session.get("role") != "admin":
            return jsonify({"success": False, "message": "投稿者本人か管理者のみピン留めできます。"}), 403
        users.pin_group_message(conn, message_id, pinned)
    finally:
        conn.close()
    return jsonify({"success": True, "message_id": message_id, "pinned": pinned})


@app.route("/group/<int:message_id>/react", methods=["POST"])
def group_message_react(message_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    emoji = request.form.get("emoji", "")
    if emoji not in REACTION_EMOJIS:
        abort(400)
    conn = _users_conn()
    try:
        if users.get_group_message(conn, message_id) is None:
            abort(404)
        existing = users.get_group_reactions_for_messages(conn, [message_id]).get(message_id, [])
        already = any(r["user_id"] == my_id and r["emoji"] == emoji for r in existing)
        if already:
            users.remove_group_reaction(conn, message_id, my_id, emoji)
        else:
            users.add_group_reaction(
                conn, message_id, my_id, emoji, now_jst_naive().strftime("%Y-%m-%d %H:%M:%S"),
            )
        reactions = users.get_group_reactions_for_messages(conn, [message_id]).get(message_id, [])
    finally:
        conn.close()
    return jsonify({"success": True, "reactions": _reaction_summary(reactions, my_id)})


@app.route("/messages")
def dm_inbox():
    """自分が参加している会話(直近やり取りが新しい順、未読件数付き)を先頭に、
    まだ会話していない他の全メンバーもそのまま同じ並び方で続ける(わざわざ相手の
    プロフィールページを経由しなくても、ここから新しい会話を始められるように
    するため)。話したことがある/ないで見た目やセクションを分けたりはしない。"""
    my_id = session.get("user_id")
    if not my_id:
        return redirect(url_for("login", next=request.path))
    conn = _users_conn()
    try:
        conversations = users.list_dm_conversations(conn, my_id)
        unread = users.dm_unread_counts(conn, my_id)
        pinned_ids = users.pinned_conversation_ids(conn, my_id)
        others = users.get_users_by_ids(conn, [c["other_id"] for c in conversations])
        all_users = users.list_users(conn)
        last_active = users.last_active_all(conn)
    finally:
        conn.close()
    entries = [
        {
            "other": others.get(c["other_id"]),
            "conversation_id": c["id"],
            "last_message_at": c["last_message_at"],
            "unread": unread.get(c["id"], 0),
            "pinned": c["id"] in pinned_ids,
            "online": _is_online(last_active.get(c["other_id"])),
        }
        for c in conversations
        if c["other_id"] in others  # 退会済みの相手は一覧に出さない
    ]
    started_ids = {c["other_id"] for c in conversations}
    entries += [
        {
            "other": u, "conversation_id": None, "last_message_at": "", "unread": 0, "pinned": False,
            "online": _is_online(last_active.get(u["id"])),
        }
        for u in all_users
        if u["id"] != my_id and u["id"] not in started_ids
    ]
    return render_template("dm_inbox.html", entries=entries)


@app.route("/messages/<int:user_id>/mark_unread", methods=["POST"])
def dm_mark_unread(user_id):
    """受信箱のクイック操作「未読にする」。まだ会話が無い相手には無意味なので、
    会話が存在する場合のみ処理する。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        a, b = sorted((my_id, user_id))
        row = conn.execute(
            "SELECT id FROM dm_conversations WHERE user_a_id = ? AND user_b_id = ?", (a, b),
        ).fetchone()
        if row is not None:
            users.mark_dm_unread(conn, row[0], my_id)
    finally:
        conn.close()
    return jsonify({"success": True})


@app.route("/messages/<int:user_id>/pin_conversation", methods=["POST"])
def dm_pin_conversation(user_id):
    """受信箱のクイック操作「ピン留め」。会話がまだ無い相手はピン留めできない
    (ピン留めは会話単位の概念のため)。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        a, b = sorted((my_id, user_id))
        row = conn.execute(
            "SELECT id FROM dm_conversations WHERE user_a_id = ? AND user_b_id = ?", (a, b),
        ).fetchone()
        if row is None:
            return jsonify({"success": False, "message": "まだ会話がありません。"}), 400
        pinned = users.toggle_pinned_conversation(conn, my_id, row[0])
    finally:
        conn.close()
    return jsonify({"success": True, "pinned": pinned})


@app.route("/messages/<int:user_id>")
def dm_thread(user_id):
    """user_idとの1対1会話を開く(存在しなければ最初のメッセージ送信時に作成されるので、
    ここでは会話が無い＝メッセージ0件として空の状態を表示するだけでよい)。"""
    my_id = session.get("user_id")
    if not my_id:
        return redirect(url_for("login", next=request.path))
    if user_id == my_id:
        abort(404)
    conn = _users_conn()
    try:
        other = users.get_user_by_id(conn, user_id)
        if other is None:
            abort(404)
        a, b = sorted((my_id, user_id))
        row = conn.execute(
            "SELECT id FROM dm_conversations WHERE user_a_id = ? AND user_b_id = ?", (a, b),
        ).fetchone()
        messages = []
        pinned = []
        other_read = 0
        if row is not None:
            conversation_id = row[0]
            messages = users.list_dm_messages(conn, conversation_id)
            message_ids = [m["id"] for m in messages]
            attachments_by_message = users.get_dm_attachments_for_messages(conn, message_ids)
            reactions_by_id = users.get_dm_reactions_for_messages(conn, message_ids)
            mentions_by_id = users.get_dm_mentions_for_messages(conn, message_ids)
            pinned = users.list_pinned_dm_messages(conn, conversation_id)
            name_by_id = {my_id: session.get("name", ""), user_id: other["name"]}
            reply_by_id = {m["id"]: m for m in messages}
            for m in messages:
                m["attachments"] = attachments_by_message.get(m["id"], [])
                m["sender_name"] = name_by_id.get(m["sender_id"], "")
            _attach_message_extras(messages, reactions_by_id, mentions_by_id, name_by_id, reply_by_id, my_id)
            for p in pinned:
                p["sender_name"] = name_by_id.get(p["sender_id"], "")
            if messages:
                users.mark_dm_read(conn, conversation_id, my_id, messages[-1]["id"])
            other_read = users.get_dm_read_state(conn, conversation_id, user_id)
        else:
            conversation_id = None
        other_online = _is_online(users.last_active(conn, user_id))
        me = users.get_user_by_id(conn, my_id)
    finally:
        conn.close()
    my_dm_theme_class = DM_THEME_BY_ID.get(me["equipped_dm_theme"], {}).get("css_class", "") if me else ""
    other_dm_theme_class = DM_THEME_BY_ID.get(other["equipped_dm_theme"], {}).get("css_class", "")
    my_bubble_style_class = BUBBLE_STYLE_BY_ID.get(me["equipped_bubble_style"], {}).get("css_class", "") if me else ""
    other_bubble_style_class = BUBBLE_STYLE_BY_ID.get(other["equipped_bubble_style"], {}).get("css_class", "")
    return render_template(
        "dm_thread.html", other=other, messages=messages, conversation_id=conversation_id,
        my_id=my_id, dm_message_max_length=DM_MESSAGE_MAX_LENGTH, other_online=other_online,
        my_dm_theme_class=my_dm_theme_class, other_dm_theme_class=other_dm_theme_class,
        my_bubble_style_class=my_bubble_style_class, other_bubble_style_class=other_bubble_style_class,
        pinned=pinned, other_read=other_read, reaction_emojis=REACTION_EMOJIS,
    )


@app.route("/messages/<int:user_id>/send", methods=["POST"])
def dm_send(user_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    if user_id == my_id:
        abort(404)
    body = request.form.get("body", "").strip()[:DM_MESSAGE_MAX_LENGTH]
    reply_to_id = request.form.get("reply_to_id", type=int)
    file = request.files.get("attachment")
    has_file = bool(file and file.filename)
    if not body and not has_file:
        return _finish(
            "dm_thread", "メッセージまたは添付ファイルを指定してください。", success=False, user_id=user_id,
        )
    attachment_info = None
    if has_file:
        result, error = _save_uploaded_file(file, DM_ATTACHMENT_DIR)
        if error:
            return _finish("dm_thread", error, success=False, user_id=user_id)
        attachment_info = result
    conn = _users_conn()
    try:
        other = users.get_user_by_id(conn, user_id)
        if other is None:
            abort(404)
        created_at = now_jst_naive().strftime("%Y-%m-%d %H:%M:%S")
        conversation_id = users.get_or_create_dm_conversation(conn, my_id, user_id, created_at)
        if reply_to_id is not None:
            reply_msg = users.get_dm_message(conn, reply_to_id)
            if reply_msg is None or reply_msg["conversation_id"] != conversation_id:
                reply_msg = None
                reply_to_id = None
        else:
            reply_msg = None
        message_id = users.send_dm_message(conn, conversation_id, my_id, body, created_at, reply_to_id=reply_to_id)
        attachment_id = None
        if attachment_info:
            original_name, stored_name, content_type, size = attachment_info
            attachment_id = users.add_dm_attachment(
                conn, message_id, original_name, stored_name, content_type, size, created_at,
            )
        name_to_id = {other["name"]: user_id}
        mention_ids = _parse_mentions(body, name_to_id)
        users.record_dm_mentions(conn, message_id, mention_ids)
        users.mark_dm_read(conn, conversation_id, my_id, message_id)
        reply_preview = _reply_preview_dict(
            reply_msg, {my_id: session.get("name", ""), user_id: other["name"]},
        )
        recipient_settings_raw = users.get_user_settings_raw(conn, user_id)
    finally:
        conn.close()
    try:
        push_notify.notify_async(
            [user_id], "dm", f"{session.get('name', '')}さんからのメッセージ",
            push_notify.message_preview(recipient_settings_raw, body, bool(attachment_info)),
            url_for("dm_thread", user_id=my_id), f"dm-{my_id}", ttl=86400, urgency="high",
        )
    except Exception as e:
        log.warning("Failed to queue the DM push notification: %s", e)
    with _typing_state_lock:
        _typing_state.pop((f"dm:{conversation_id}", my_id), None)
    if _wants_json():
        return jsonify({
            "message": "送信しました。", "success": True,
            "id": message_id, "body": body, "created_at": created_at,
            "attachment": (
                {"id": attachment_id, "filename": attachment_info[0]} if attachment_info else None
            ),
            "reply_preview": reply_preview,
        })
    return redirect(url_for("dm_thread", user_id=user_id))


@app.route("/messages/<int:user_id>/poll")
def dm_poll(user_id):
    """開いているスレッドで新着メッセージをJSで取りに行くためのエンドポイント
    (このアプリには汎用のリアルタイム配信基盤が無いため、スレッド表示中だけ
    数秒おきにポーリングする軽量な方式にしている)。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    after_id = request.args.get("after", 0, type=int)
    conn = _users_conn()
    try:
        other = users.get_user_by_id(conn, user_id)
        a, b = sorted((my_id, user_id))
        row = conn.execute(
            "SELECT id FROM dm_conversations WHERE user_a_id = ? AND user_b_id = ?", (a, b),
        ).fetchone()
        messages = []
        other_read = 0
        conversation_id = row[0] if row is not None else None
        if row is not None:
            messages = users.list_dm_messages(conn, conversation_id, after_id)
            message_ids = [m["id"] for m in messages]
            attachments_by_message = users.get_dm_attachments_for_messages(conn, message_ids)
            reactions_by_id = users.get_dm_reactions_for_messages(conn, message_ids)
            mentions_by_id = users.get_dm_mentions_for_messages(conn, message_ids)
            name_by_id = {my_id: session.get("name", ""), user_id: other["name"] if other else ""}
            reply_ids = {m["reply_to_id"] for m in messages if m.get("reply_to_id")}
            reply_msgs_by_id = {rid: users.get_dm_message(conn, rid) for rid in reply_ids}
            for m in messages:
                m["attachments"] = attachments_by_message.get(m["id"], [])
                m["sender_name"] = name_by_id.get(m["sender_id"], "")
                m["reactions"] = _reaction_summary(reactions_by_id.get(m["id"], []), my_id)
                mention_ids = mentions_by_id.get(m["id"], [])
                m["mentioned_names"] = [name_by_id[uid] for uid in mention_ids if uid in name_by_id]
                try:
                    m["card"] = json.loads(m["card_data"]) if m.get("card_data") else None
                except (TypeError, ValueError):
                    m["card"] = None
                m["reply_preview"] = _reply_preview_dict(reply_msgs_by_id.get(m.get("reply_to_id")), name_by_id)
            if messages:
                users.mark_dm_read(conn, conversation_id, my_id, messages[-1]["id"])
            other_read = users.get_dm_read_state(conn, conversation_id, user_id)
        other_online = _is_online(users.last_active(conn, user_id))
    finally:
        conn.close()
    typing_names = []
    if conversation_id is not None and other is not None:
        typing_ids = _typing_user_ids(f"dm:{conversation_id}", exclude_user_id=my_id)
        if user_id in typing_ids:
            typing_names = [other["name"]]
    return jsonify({
        "messages": messages, "other_online": other_online, "other_read": other_read,
        "typing_names": typing_names,
    })


@app.route("/messages/attachments/<int:attachment_id>")
def dm_attachment_file(attachment_id):
    """DM添付ファイルのダウンロード。その会話の当事者(送信者・受信者どちらか)だけが
    取得できる(会話に無関係な人がidを推測してもファイルを見られないようにするため)。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        att = users.get_dm_attachment(conn, attachment_id)
    finally:
        conn.close()
    if att is None:
        abort(404)
    if my_id not in (att["user_a_id"], att["user_b_id"]):
        abort(403)
    return send_from_directory(
        DM_ATTACHMENT_DIR, att["stored_name"], as_attachment=True, download_name=att["filename"],
    )


@app.route("/messages/<int:user_id>/<int:message_id>/delete", methods=["POST"])
def dm_message_delete(user_id, message_id):
    """自分が送信したメッセージを削除する(相手のメッセージは消せない。右クリック
    メニューからJS経由で叩かれる想定なので、成功/失敗をJSONで返すだけでよい)。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        msg = users.get_dm_message(conn, message_id)
        if msg is None or my_id not in (msg["user_a_id"], msg["user_b_id"]):
            abort(404)
        if msg["sender_id"] != my_id:
            return jsonify({"success": False, "message": "自分が送信したメッセージのみ削除できます。"}), 403
        stored_names = users.get_dm_message_attachment_stored_names(conn, message_id)
        users.delete_dm_message(conn, message_id)
    finally:
        conn.close()
    for stored_name in stored_names:
        (DM_ATTACHMENT_DIR / stored_name).unlink(missing_ok=True)
    return jsonify({"success": True, "message_id": message_id})


@app.route("/messages/<int:user_id>/<int:message_id>/edit", methods=["POST"])
def dm_message_edit(user_id, message_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    body = request.form.get("body", "").strip()[:DM_MESSAGE_MAX_LENGTH]
    if not body:
        return jsonify({"success": False, "message": "本文を入力してください。"}), 400
    conn = _users_conn()
    try:
        msg = users.get_dm_message(conn, message_id)
        if msg is None or my_id not in (msg["user_a_id"], msg["user_b_id"]):
            abort(404)
        if msg["sender_id"] != my_id:
            return jsonify({"success": False, "message": "自分が送信したメッセージのみ編集できます。"}), 403
        edited_at = now_jst_naive().strftime("%Y-%m-%d %H:%M:%S")
        users.edit_dm_message(conn, message_id, body, edited_at)
    finally:
        conn.close()
    return jsonify({"success": True, "message_id": message_id, "body": body, "edited_at": edited_at})


@app.route("/messages/<int:user_id>/<int:message_id>/pin", methods=["POST"])
def dm_message_pin(user_id, message_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    pinned = request.form.get("pinned") == "1"
    conn = _users_conn()
    try:
        msg = users.get_dm_message(conn, message_id)
        if msg is None or my_id not in (msg["user_a_id"], msg["user_b_id"]):
            abort(404)
        users.pin_dm_message(conn, message_id, pinned)
    finally:
        conn.close()
    return jsonify({"success": True, "message_id": message_id, "pinned": pinned})


@app.route("/messages/<int:user_id>/<int:message_id>/react", methods=["POST"])
def dm_message_react(user_id, message_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    emoji = request.form.get("emoji", "")
    if emoji not in REACTION_EMOJIS:
        abort(400)
    conn = _users_conn()
    try:
        msg = users.get_dm_message(conn, message_id)
        if msg is None or my_id not in (msg["user_a_id"], msg["user_b_id"]):
            abort(404)
        existing = users.get_dm_reactions_for_messages(conn, [message_id]).get(message_id, [])
        already = any(r["user_id"] == my_id and r["emoji"] == emoji for r in existing)
        if already:
            users.remove_dm_reaction(conn, message_id, my_id, emoji)
        else:
            users.add_dm_reaction(conn, message_id, my_id, emoji, now_jst_naive().strftime("%Y-%m-%d %H:%M:%S"))
        reactions = users.get_dm_reactions_for_messages(conn, [message_id]).get(message_id, [])
    finally:
        conn.close()
    return jsonify({"success": True, "reactions": _reaction_summary(reactions, my_id)})


@app.route("/messages/<int:user_id>/typing", methods=["POST"])
def dm_typing(user_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    a, b = sorted((my_id, user_id))
    conn = _users_conn()
    try:
        row = conn.execute(
            "SELECT id FROM dm_conversations WHERE user_a_id = ? AND user_b_id = ?", (a, b),
        ).fetchone()
        if row is None:
            created_at = now_jst_naive().strftime("%Y-%m-%d %H:%M:%S")
            conversation_id = users.get_or_create_dm_conversation(conn, my_id, user_id, created_at)
        else:
            conversation_id = row[0]
    finally:
        conn.close()
    _mark_typing(f"dm:{conversation_id}", my_id)
    return jsonify({"success": True})


@app.route("/messages/<int:user_id>/search")
def dm_search(user_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    keyword = request.args.get("q", "").strip()
    if not keyword:
        return jsonify({"results": []})
    a, b = sorted((my_id, user_id))
    conn = _users_conn()
    try:
        row = conn.execute(
            "SELECT id FROM dm_conversations WHERE user_a_id = ? AND user_b_id = ?", (a, b),
        ).fetchone()
        if row is None:
            return jsonify({"results": []})
        matches = users.search_dm_messages(conn, row[0], keyword, limit=MESSAGE_SEARCH_LIMIT)
    finally:
        conn.close()
    return jsonify({"results": matches})


_PREVIEWABLE_XLSX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_PREVIEW_MAX_ROWS = 200
_PREVIEW_MAX_COLS = 30

# /files/<id>/raw はブラウザにその場で描画させる生配信のため、content_typeが
# text/html等になりうる場合はここで弾かないと格納型XSSが成立してしまう。
# アップロード自体には拡張子制限を設けていない(利用者の要望で撤廃済み)ため、
# 「その場で描画してよい」ものだけをここで安全側に絞り、それ以外は
# ダウンロード(添付ヘッダー付き)にフォールバックする。プレビュー画面が
# <img>/<embed>で読み込むのは画像とPDFのみ(shared_file_preview.html参照)。
# その場表示(添付ヘッダーなし)を許す種類。以前は"image/"で始まれば何でも通していたため、
# SVG(image/svg+xml、中にスクリプトを書ける)をアップロードして/files/<id>/rawを開かせると
# このサイトの権限でスクリプトが動いてしまった(2026-10-08の監査で指摘)。
SAFE_INLINE_CONTENT_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp", "application/pdf"}


def _is_safe_inline_content_type(content_type):
    return (content_type or "").split(";")[0].strip().lower() in SAFE_INLINE_CONTENT_TYPES


@app.route("/files")
def shared_files_page():
    """班の共有ファイル置き場(誰でもアップロード・ダウンロードできる、Teamsの「ファイル」
    タブに相当)。ログイン済みなら誰でも見られる。?folder=<id>でフォルダの中を表示する
    (省略時は最上位)。?q=キーワードを指定すると、フォルダをまたいで全文検索する
    (ファイル名・説明・自動タグ・テキスト系ファイルの中身が対象)。?starred=1で
    お気に入りのみに絞り込める(検索と併用可能)。"""
    if not session.get("user_id"):
        return redirect(url_for("login", next=request.path))
    my_id = session["user_id"]
    folder_id = request.args.get("folder", type=int)
    query = request.args.get("q", "").strip()
    starred_only = request.args.get("starred") == "1"
    conn = _users_conn()
    try:
        starred_ids = users.starred_file_ids(conn, my_id)
        if query:
            folders = []
            files = users.search_shared_files(conn, query)
            breadcrumbs = []
        else:
            if folder_id is not None and users.get_shared_folder(conn, folder_id) is None:
                abort(404)
            folders = users.list_shared_folders(conn, folder_id)
            files = users.list_shared_files(conn, folder_id)
            breadcrumbs = users.shared_folder_breadcrumbs(conn, folder_id) if folder_id else []
        if starred_only:
            files = [f for f in files if f["id"] in starred_ids]
        creators = users.get_users_by_ids(
            conn, [f["uploaded_by"] for f in files] + [fo["created_by"] for fo in folders]
        )
    finally:
        conn.close()
    for f in files:
        f["uploader_name"] = creators.get(f["uploaded_by"], {}).get("name", "(退会済みユーザー)")
        f["previewable"] = (
            f["content_type"].startswith("image/") or f["content_type"] == "application/pdf"
            or f["content_type"] in _PREVIEWABLE_XLSX_CONTENT_TYPES
        )
        f["tag_list"] = [t for t in f["tags"].split(",") if t]
        f["starred"] = f["id"] in starred_ids
    for fo in folders:
        fo["created_by_name"] = creators.get(fo["created_by"], {}).get("name", "(退会済みユーザー)")
    return render_template(
        "shared_files.html", files=files, folders=folders, breadcrumbs=breadcrumbs,
        current_folder_id=folder_id, query=query, starred_only=starred_only,
    )


_TEXT_INDEXABLE_EXTENSIONS = {
    ".txt": "text/plain", ".log": "text/plain", ".csv": "text/csv",
    ".md": "text/markdown", ".json": "application/json",
}
CONTENT_TEXT_MAX_BYTES = 300 * 1024
_DATE_TAG_RE = re.compile(r"(\d{4})[-_]?(\d{2})[-_]?(\d{2})")


def _auto_tags_for_filename(filename):
    """アップロードされたファイル名から、地点名(STATION_NAMESとの部分一致)・日付らしき
    文字列を自動で拾い、"地点:丹過橋,日付:2026-09-03"のようなタグ文字列を組み立てる。
    捏造は避け、あくまでファイル名自体に既に含まれている情報を拾うだけにとどめる。"""
    tags = [f"地点:{name}" for name in STATION_NAMES if name and name in filename]
    m = _DATE_TAG_RE.search(filename)
    if m:
        tags.append(f"日付:{m.group(1)}-{m.group(2)}-{m.group(3)}")
    return ",".join(tags)


def _extract_content_text(path, ext):
    """検索(全文検索)用に、テキスト系ファイルの中身を読み込む。対象拡張子以外は
    空文字を返す(PDF/画像/xlsx等の中身までは読まない。新規の重い依存ライブラリを
    増やさない方針のため)。"""
    if ext not in _TEXT_INDEXABLE_EXTENSIONS:
        return ""
    try:
        with open(path, "rb") as f:
            raw = f.read(CONTENT_TEXT_MAX_BYTES)
        return raw.decode("utf-8", errors="ignore")
    except OSError:
        return ""


@app.route("/files/upload", methods=["POST"])
def shared_file_upload():
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    folder_id = request.form.get("folder_id", type=int)
    file = request.files.get("file")
    if not file or not file.filename:
        return _finish("shared_files_page", "ファイルを選択してください。", success=False, folder=folder_id)
    result, error = _save_uploaded_file(file, SHARED_FILES_DIR)
    if error:
        return _finish("shared_files_page", error, success=False, folder=folder_id)
    original_name, stored_name, content_type, size = result
    description = request.form.get("description", "").strip()[:200]
    tags = _auto_tags_for_filename(original_name)
    content_text = _extract_content_text(SHARED_FILES_DIR / stored_name, Path(original_name).suffix.lower())
    conn = _users_conn()
    try:
        if folder_id is not None and users.get_shared_folder(conn, folder_id) is None:
            abort(404)
        users.add_shared_file(
            conn, original_name, stored_name, content_type, size, my_id,
            now_jst_naive().strftime("%Y-%m-%d %H:%M:%S"), description, folder_id,
            tags=tags, content_text=content_text,
        )
    finally:
        conn.close()
    return _finish("shared_files_page", "アップロードしました。", folder=folder_id)


@app.route("/files/folders", methods=["POST"])
def shared_folder_create():
    """右クリックメニューの「新しいフォルダ」から呼ばれる。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    parent_id = request.form.get("parent_id", type=int)
    name = request.form.get("name", "").strip()[:100]
    if not name:
        return _finish("shared_files_page", "フォルダ名を入力してください。", success=False, folder=parent_id)
    conn = _users_conn()
    try:
        if parent_id is not None and users.get_shared_folder(conn, parent_id) is None:
            abort(404)
        users.create_shared_folder(conn, name, parent_id, my_id, now_jst_naive().strftime("%Y-%m-%d %H:%M:%S"))
    finally:
        conn.close()
    return _finish("shared_files_page", "フォルダを作成しました。", folder=parent_id)


@app.route("/files/folders/<int:folder_id>/delete", methods=["POST"])
def shared_folder_delete(folder_id):
    """空のフォルダのみ削除できる(中身を巻き添えで消さないため)。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        folder = users.get_shared_folder(conn, folder_id)
        if folder is None:
            abort(404)
        if folder["created_by"] != my_id and session.get("role") != "admin":
            abort(403)
        # 空チェックと削除はusers.delete_shared_folder内で同じトランザクションとして行う
        # (以前はここで別々に呼んでおり、チェックと削除の間に別リクエストがファイルを
        # 追加するレースで未処理のIntegrityErrorになりうった、2026-09-13の監査で指摘)。
        if not users.delete_shared_folder(conn, folder_id):
            return _finish(
                "shared_files_page", "空でないフォルダは削除できません。",
                success=False, folder=folder["parent_id"],
            )
        parent_id = folder["parent_id"]
    finally:
        conn.close()
    return _finish("shared_files_page", "フォルダを削除しました。", folder=parent_id)


@app.route("/files/<int:file_id>")
def shared_file_download(file_id):
    if not session.get("user_id"):
        abort(403)
    conn = _users_conn()
    try:
        f = users.get_shared_file(conn, file_id)
    finally:
        conn.close()
    if f is None:
        abort(404)
    return send_from_directory(
        SHARED_FILES_DIR, f["stored_name"], as_attachment=True, download_name=f["filename"],
    )


@app.route("/files/<int:file_id>/raw")
def shared_file_raw(file_id):
    """プレビュー画面の<img>/<embed>から読み込ませる、添付ヘッダーを付けない生配信。"""
    if not session.get("user_id"):
        abort(403)
    conn = _users_conn()
    try:
        f = users.get_shared_file(conn, file_id)
    finally:
        conn.close()
    if f is None:
        abort(404)
    if not _is_safe_inline_content_type(f["content_type"]):
        # 画像・PDF以外をその場描画させるとcontent_type偽装によるXSSが成立するため、
        # 添付ダウンロードにフォールバックする(アップロード自体は拒否しない)。
        return send_from_directory(
            SHARED_FILES_DIR, f["stored_name"], as_attachment=True, download_name=f["filename"],
        )
    return send_from_directory(SHARED_FILES_DIR, f["stored_name"], mimetype=f["content_type"])


@app.route("/files/<int:file_id>/preview")
def shared_file_preview(file_id):
    """画像・PDF・xlsxをブラウザ上でその場に見られるプレビュー画面。
    xlsxはopenpyxlで読み込み、先頭シートを表としてサーバー側でHTML化する
    (行・列数が多いシートは重くなるので上限で打ち切る)。"""
    if not session.get("user_id"):
        abort(403)
    conn = _users_conn()
    try:
        f = users.get_shared_file(conn, file_id)
    finally:
        conn.close()
    if f is None:
        abort(404)

    kind = None
    xlsx_rows = None
    xlsx_truncated = False
    if f["content_type"].startswith("image/"):
        kind = "image"
    elif f["content_type"] == "application/pdf":
        kind = "pdf"
    elif f["content_type"] in _PREVIEWABLE_XLSX_CONTENT_TYPES:
        kind = "xlsx"
        import openpyxl
        try:
            wb = openpyxl.load_workbook(
                SHARED_FILES_DIR / f["stored_name"], data_only=True, read_only=True,
            )
            ws = wb.active
            xlsx_rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _PREVIEW_MAX_ROWS:
                    xlsx_truncated = True
                    break
                xlsx_rows.append(row[:_PREVIEW_MAX_COLS])
            if ws.max_column and ws.max_column > _PREVIEW_MAX_COLS:
                xlsx_truncated = True
        except Exception as e:
            log.warning("xlsx preview failed for shared_files id=%d: %s", file_id, e)
            kind = None
    return render_template(
        "shared_file_preview.html", f=f, kind=kind, xlsx_rows=xlsx_rows, xlsx_truncated=xlsx_truncated,
    )


@app.route("/files/<int:file_id>/delete", methods=["POST"])
def shared_file_delete(file_id):
    """アップロード本人か管理者だけが削除できる。"""
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        f = users.get_shared_file(conn, file_id)
        if f is None:
            abort(404)
        if f["uploaded_by"] != my_id and session.get("role") != "admin":
            abort(403)
        users.delete_shared_file(conn, file_id)
        folder_id = f["folder_id"]
    finally:
        conn.close()
    stored_path = SHARED_FILES_DIR / f["stored_name"]
    stored_path.unlink(missing_ok=True)
    return _finish("shared_files_page", "削除しました。", folder=folder_id)


@app.route("/files/<int:file_id>/star", methods=["POST"])
def shared_file_star(file_id):
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    conn = _users_conn()
    try:
        f = users.get_shared_file(conn, file_id)
        if f is None:
            abort(404)
        starred = users.toggle_file_star(conn, my_id, file_id, now_jst_naive().strftime("%Y-%m-%d %H:%M:%S"))
    finally:
        conn.close()
    if _wants_json():
        return jsonify({"success": True, "starred": starred})
    return _finish("shared_files_page", "お気に入りに追加しました。" if starred else "お気に入りを解除しました。")


def _bulk_file_ids_from_form():
    ids = request.form.getlist("file_ids", type=int)
    return sorted(set(i for i in ids if i))


@app.route("/files/bulk/move", methods=["POST"])
def shared_file_bulk_move():
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    file_ids = _bulk_file_ids_from_form()
    target_folder_id = request.form.get("target_folder_id", type=int)
    if not file_ids:
        return jsonify({"success": False, "message": "移動するファイルを選択してください。"}), 400
    conn = _users_conn()
    try:
        if target_folder_id is not None and users.get_shared_folder(conn, target_folder_id) is None:
            return jsonify({"success": False, "message": "移動先フォルダが見つかりません。"}), 404
        moved = 0
        for file_id in file_ids:
            f = users.get_shared_file(conn, file_id)
            if f is None:
                continue
            if f["uploaded_by"] != my_id and session.get("role") != "admin":
                continue
            users.move_shared_file(conn, file_id, target_folder_id)
            moved += 1
    finally:
        conn.close()
    skipped = len(file_ids) - moved
    message = f"{moved}件移動しました。" + (f"({skipped}件は権限がなくスキップしました)" if skipped else "")
    return jsonify({"success": True, "message": message})


@app.route("/files/bulk/delete", methods=["POST"])
def shared_file_bulk_delete():
    my_id = session.get("user_id")
    if not my_id:
        abort(403)
    file_ids = _bulk_file_ids_from_form()
    if not file_ids:
        return jsonify({"success": False, "message": "削除するファイルを選択してください。"}), 400
    conn = _users_conn()
    try:
        deleted = 0
        stored_paths = []
        for file_id in file_ids:
            f = users.get_shared_file(conn, file_id)
            if f is None:
                continue
            if f["uploaded_by"] != my_id and session.get("role") != "admin":
                continue
            users.delete_shared_file(conn, file_id)
            stored_paths.append(SHARED_FILES_DIR / f["stored_name"])
            deleted += 1
    finally:
        conn.close()
    for p in stored_paths:
        p.unlink(missing_ok=True)
    skipped = len(file_ids) - deleted
    message = f"{deleted}件削除しました。" + (f"({skipped}件は権限がなくスキップしました)" if skipped else "")
    return jsonify({"success": True, "message": message})


@app.route("/files/bulk/download")
def shared_file_bulk_download():
    """複数選択したファイルをまとめてZIPでダウンロードする。build_data_backup_zip
    (バックアップ機能)と同じzipfile標準ライブラリを使う。"""
    if not session.get("user_id"):
        abort(403)
    file_ids = sorted(set(request.args.getlist("file_ids", type=int)))
    if not file_ids:
        abort(400)
    conn = _users_conn()
    try:
        files = [users.get_shared_file(conn, fid) for fid in file_ids]
    finally:
        conn.close()
    files = [f for f in files if f is not None]
    if not files:
        abort(404)
    buf = io.BytesIO()
    used_names = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in files:
            path = SHARED_FILES_DIR / f["stored_name"]
            if not path.exists():
                continue
            arcname = f["filename"]
            # 同名ファイルが複数選択された場合、ZIP内で上書きされないように連番を振る。
            n = 1
            base_arcname = arcname
            while arcname in used_names:
                stem, dot, ext = base_arcname.rpartition(".")
                arcname = f"{stem or base_arcname}_{n}{dot}{ext}" if dot else f"{base_arcname}_{n}"
                n += 1
            used_names.add(arcname)
            zf.write(path, arcname=arcname)
    buf.seek(0)
    return Response(
        buf.getvalue(), mimetype="application/zip",
        headers={"Content-Disposition": "attachment; filename=shared_files.zip"},
    )
