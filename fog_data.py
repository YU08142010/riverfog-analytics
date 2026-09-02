"""只見川.xlsxを読み込み、霧予測モデル用のDataFrameに変換する。

common.pyの固定列レイアウトを前提とする（列自動検出は行わない）。
"""
import csv
import threading
from datetime import datetime

import openpyxl
import pandas as pd

from common import (
    CAMERAS_CSV,
    COL_DATETIME,
    COL_DEWPOINT,
    COL_HUMIDITY,
    COL_PHENOM_START,
    COL_PRECIP,
    COL_SOLAR,
    COL_SUNSHINE,
    COL_TEMP,
    COL_VISIBILITY,
    COL_WIND_DIR,
    COL_WIND_SPEED,
    DATA_XLSX,
)
from fog_features import (
    COL_DEWPOINT as F_DEWPOINT,
    COL_HUMIDITY as F_HUMIDITY,
    COL_PRECIP as F_PRECIP,
    COL_SOLAR as F_SOLAR,
    COL_SUNSHINE as F_SUNSHINE,
    COL_TEMP as F_TEMP,
    COL_VISIBILITY as F_VISIBILITY,
    COL_WIND_DIR_DEG as F_WIND_DIR_DEG,
    COL_WIND_SPEED as F_WIND_SPEED,
    encode_wind_direction_degrees,
)


def load_station_names(cameras_csv=CAMERAS_CSV):
    with open(cameras_csv, encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        return [row["地点名"] for row in reader]


# openpyxlのread_onlyモードでもiter_rows(values_only=True)自体の解析コストは行×列数に
# 比例して重く、データが数年分溜まった現状(19000行超)では1回の読み込みに10秒以上かかる
# （ws.cell()のセル単位ランダムアクセスほど酷くはないが、無視できる差ではない）。
# webapp.py側のほぼ全ページがリクエストのたびにこの関数を呼ぶため、キャッシュ無しだと
# サイト全体が常にこの待ち時間を背負うことになる。xlsxのファイル更新時刻(mtime)をキーに
# キャッシュし、collect.py等が実際に書き込むまでは再パースしない。
_cache_lock = threading.Lock()
_cache = {"xlsx_path": None, "mtime": None, "data": None}


def load_river_fog_data(xlsx_path=DATA_XLSX, cameras_csv=CAMERAS_CSV):
    """(main_df, phenom_df, station_names) を返す。呼び出し元ごとに独立した
    DataFrameのコピーを返すので、呼び出し側で自由に加工してよい(キャッシュ本体は
    汚染されない)。

    main_df  : datetime + 気象特徴量（fog_features.py の列名）
    phenom_df: datetime + 地点名ごとの現象コード生値（"/" / 数値 / None）
    """
    try:
        mtime = xlsx_path.stat().st_mtime
    except OSError:
        mtime = None
    with _cache_lock:
        if _cache["xlsx_path"] == xlsx_path and _cache["mtime"] == mtime and _cache["data"] is not None:
            main_df, phenom_df, station_names = _cache["data"]
            return main_df.copy(), phenom_df.copy(), list(station_names)

    main_df, phenom_df, station_names = _load_river_fog_data_uncached(xlsx_path, cameras_csv)

    with _cache_lock:
        _cache["xlsx_path"] = xlsx_path
        _cache["mtime"] = mtime
        _cache["data"] = (main_df, phenom_df, station_names)
    return main_df.copy(), phenom_df.copy(), list(station_names)


def _load_river_fog_data_uncached(xlsx_path, cameras_csv):
    """load_river_fog_data()の実処理(キャッシュ未ヒット時のみ呼ばれる)。

    openpyxlのread_onlyモードでは ws.cell(row, col) によるセル単位のランダムアクセスは
    行ごとに内部XMLの再走査が発生し、行数に対してほぼ二乗オーダーで遅くなる
    （データ件数がまだ少ない段階でも数分かかることを確認済み）。
    必ず iter_rows(values_only=True) で行単位にまとめて取得すること。
    """
    station_names = load_station_names(cameras_csv)
    n_cols = COL_PHENOM_START - 1 + len(station_names)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.active

    main_rows = []
    phenom_rows = []
    for values in ws.iter_rows(min_row=1, max_col=n_cols, values_only=True):
        dt = values[COL_DATETIME - 1]
        if not isinstance(dt, datetime):
            continue  # 見出し行・小見出し行をスキップ

        main_rows.append({
            "datetime": dt,
            F_TEMP: values[COL_TEMP - 1],
            F_PRECIP: values[COL_PRECIP - 1],
            F_WIND_SPEED: values[COL_WIND_SPEED - 1],
            F_WIND_DIR_DEG: encode_wind_direction_degrees(values[COL_WIND_DIR - 1]),
            F_SOLAR: values[COL_SOLAR - 1],
            F_SUNSHINE: values[COL_SUNSHINE - 1],
            F_DEWPOINT: values[COL_DEWPOINT - 1],
            F_HUMIDITY: values[COL_HUMIDITY - 1],
            F_VISIBILITY: values[COL_VISIBILITY - 1],
        })

        phenom_row = {"datetime": dt}
        for i, name in enumerate(station_names):
            phenom_row[name] = values[COL_PHENOM_START - 1 + i]
        phenom_rows.append(phenom_row)
    wb.close()

    main_df = pd.DataFrame(main_rows)
    for col in (F_TEMP, F_PRECIP, F_WIND_SPEED, F_SOLAR, F_SUNSHINE, F_DEWPOINT, F_HUMIDITY, F_VISIBILITY):
        main_df[col] = pd.to_numeric(main_df[col], errors="coerce")
    phenom_df = pd.DataFrame(phenom_rows)
    return main_df, phenom_df, station_names
