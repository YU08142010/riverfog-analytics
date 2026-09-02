"""欠測しているカメラ画像を、cam.river.go.jpのhistoryエンドポイントから遡って取得し直す
バックフィルスクリプト（気象データのbackfill_weather.pyと同じ思想）。

collect.pyの毎時取得(now.jpg)は、DNS障害やカメラサイト側の混雑等で一時的に失敗することがある。
失敗時は data/images/{date}/{HHMM}/{地点名}_unavailable.jpg というプレースホルダー画像が
保存される（Geminiが「実際の写真ではない」と判定した場合も同じ命名）。

本スクリプトはそのプレースホルダーを手がかりに、
cam.river.go.jp/cam/history/{YYYYMMDDHHMM}/{camId}.jpg から該当時刻のスナップショットを
取得し直す。history側は5分刻みでしか保存されておらず、ちょうどその時刻のスナップショットも
無いことがある（サイト側で元々撮れていなかった場合）ため、前後 HISTORY_SEARCH_WINDOW_MIN 分の
範囲を5分刻みで近い順に探す（河川・霧の状況は数分では大きく変わらないため、
近い時刻のスナップショットでも実用上十分）。

取得できた画像は、collect.pyと同じGeminiモデル・プロンプトで分類し直し、xlsxの該当セルと
confidence_log.csvを埋める。Geminiが「実際の写真ではない」（-1）と判定した場合は
プレースホルダーのまま維持する（元々写真として成立していなかった可能性が高いため）。

collect.pyと同時に実行すると画像アーカイブ・confidence_log.csv・xlsxの書き込みが競合するため、
毎時収集(04:00-19:00台の00分)を避けたタイミングでcron等から定期実行すること
（backfill_weather.pyと同じ注意点）。また同じxlsxファイルをbackfill_weather.pyも書き換えるため、
両者が同時に走らないよう実行時刻をずらすこと（crontabではbackfill_weather.py 19:20、
predict.py 19:30の後、本スクリプトは19:45に設定している）。
"""
import logging
import logging.handlers
import os
import sys
from datetime import datetime, timedelta

import openpyxl
import requests
from dotenv import load_dotenv

import collect
from common import (
    BASE_DIR,
    COL_DATETIME,
    COL_PHENOM_START,
    DATA_XLSX,
    DEFAULT_GEMINI_MODEL,
    JST,
    XlsxLockTimeout,
    record_network_failure,
    request_with_dns_fallback,
    xlsx_write_lock,
)

HISTORY_URL = "https://cam.river.go.jp/cam/history/{ts}/{cam_id}.jpg"
HISTORY_SEARCH_STEP_MIN = 5  # history側の実際の保存間隔（5分刻み以外は404になる）
HISTORY_SEARCH_WINDOW_MIN = 30  # 目標時刻の前後何分まで探すか
HISTORY_FETCH_TIMEOUT = 10

# history側の保持期間は実測で28日前(成功)〜2ヶ月前(404)の間だったため、安全側に見て
# これより古い欠測は個別ログを出さず件数のみ集計する
GIVE_UP_AFTER_DAYS = 21

LOG_FILE = BASE_DIR / "logs" / "backfill_images.log"
LOG_MAX_BYTES = 2 * 1024 * 1024
LOG_BACKUP_COUNT = 3

# collect.pyをimportすると、そちら側でlogging.basicConfig()が既に実行済みになる
# （最初の1回しか効かないため、ここでbasicConfig()しても無視される）。そのため
# ルートロガーには頼らず、このロガー専用のハンドラを直接持たせて完全に分離する
# （さもないとbackfill_images.log向けのログがcollect.log側に混ざってしまう）。
log = logging.getLogger("backfill_images")
log.setLevel(logging.INFO)
log.propagate = False
_formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
_file_handler = logging.handlers.RotatingFileHandler(
    LOG_FILE, maxBytes=LOG_MAX_BYTES, backupCount=LOG_BACKUP_COUNT, encoding="utf-8"
)
_file_handler.setFormatter(_formatter)
_stream_handler = logging.StreamHandler()
_stream_handler.setFormatter(_formatter)
log.addHandler(_file_handler)
log.addHandler(_stream_handler)


def _search_offsets():
    """0, +5, -5, +10, -10, ... の順で、目標時刻に近いオフセット(分)を返す。"""
    yield 0
    step = HISTORY_SEARCH_STEP_MIN
    while step <= HISTORY_SEARCH_WINDOW_MIN:
        yield step
        yield -step
        step += HISTORY_SEARCH_STEP_MIN


def fetch_history_image(cam_id, target_dt, station_name=None):
    """target_dtちょうど、またはその前後の近いスナップショットを取得する。
    成功したら (画像bytes, 実際に取れた時刻) を返す。見つからなければ (None, None)。
    """
    connection_failed = False
    for offset in _search_offsets():
        dt = target_dt + timedelta(minutes=offset)
        ts = dt.strftime("%Y%m%d%H%M")
        url = HISTORY_URL.format(ts=ts, cam_id=cam_id)
        try:
            resp = request_with_dns_fallback(
                "GET", url, headers=collect.CAM_HEADERS, timeout=HISTORY_FETCH_TIMEOUT, log=log,
            )
        except requests.RequestException:
            # DNS障害時は全オフセットで同じ理由で失敗するため、記録は1回にまとめる
            connection_failed = True
            continue
        if resp.status_code == 200 and resp.content:
            return resp.content, dt
    if connection_failed:
        record_network_failure(
            "backfill_images", f"cam_id={cam_id}, target={target_dt}", station=station_name,
        )
    return None, None


# collect.pyは"_unavailable.jpg"を2種類の理由で書く: (1)取得自体に失敗した場合の
# 灰色プレースホルダー画像、(2)取得・分類には成功したが、Geminiが「実際の写真ではない」
# (-1)と判定した実写真そのもの。(2)は既に実物の写真をGeminiが見た上での判定であり、
# history APIから同時刻の別スナップショットを取り直しても同じ判定になる可能性が高いため、
# (1)の灰色プレースホルダーだけを再取得対象にする(でないと無駄なAPI呼び出しが延々続く)。
# 判別は、collect.pyが書く灰色プレースホルダーの中身が常に同一バイト列であることを利用する。
_FETCH_FAILED_PLACEHOLDER_BYTES = collect._fetch_failed_placeholder_bytes()


def find_missing_slots():
    """data/images配下の_unavailable.jpgのうち、取得自体に失敗した(灰色プレースホルダーの)
    ものだけを走査し、(date_str, hhmm, station_name)のリストを返す。
    観測時間帯(4-19時)のみを対象にする。
    """
    slots = []
    if not collect.IMAGE_ARCHIVE_DIR.exists():
        return slots
    for day_dir in sorted(collect.IMAGE_ARCHIVE_DIR.iterdir()):
        if not day_dir.is_dir():
            continue
        try:
            datetime.strptime(day_dir.name, "%Y-%m-%d")
        except ValueError:
            continue
        for hour_dir in sorted(day_dir.iterdir()):
            if not hour_dir.is_dir() or len(hour_dir.name) != 4:
                continue
            hour = int(hour_dir.name[:2])
            if not (collect.OBSERVATION_START_HOUR <= hour <= collect.OBSERVATION_END_HOUR):
                continue
            for f in hour_dir.glob("*_unavailable.jpg"):
                try:
                    is_fetch_failure = f.read_bytes() == _FETCH_FAILED_PLACEHOLDER_BYTES
                except OSError:
                    continue
                if not is_fetch_failure:
                    continue  # Gemini判定済みの実写真。再取得しても無駄なのでスキップ
                name = f.name[: -len("_unavailable.jpg")]
                slots.append((day_dir.name, hour_dir.name, name))
    return slots


def _row_index_map(ws):
    """datetime -> 行番号 の対応表を作る（全行を一度だけ走査）。"""
    idx = {}
    for row in range(2, ws.max_row + 1):
        dt = ws.cell(row=row, column=COL_DATETIME).value
        if isinstance(dt, datetime):
            idx[dt] = row
    return idx


def main():
    load_dotenv(BASE_DIR / ".env")
    dry_run = "--dry-run" in sys.argv[1:]
    log.info("=== Backfill started%s ===", " (dry-run)" if dry_run else "")

    # xlsxへの書き込み(load_workbook〜save)はcollect.py・backfill_weather.pyとも競合しうるため
    # 排他ロックを取る(common.xlsx_write_lock参照)。この関数は画像取得・Gemini分類など
    # xlsxに触れない処理も多いが、途中で書き込みが発生するかどうかは事前に分からないため、
    # 実行全体を通してロックを保持する(バックグラウンドバッチ処理なので多少長く保持しても
    # 実害は小さいと判断)。
    try:
        with xlsx_write_lock():
            _run_backfill(dry_run)
    except XlsxLockTimeout as e:
        log.error("%s; skipping this run (collect.py or backfill_weather.py is likely writing)", e)


def _run_backfill(dry_run):
    cameras = collect.load_cameras()
    cam_id_by_name = dict(cameras)
    col_by_name = {name: COL_PHENOM_START + i for i, (name, _cam_id) in enumerate(cameras)}

    slots = find_missing_slots()
    if not slots:
        log.info("No missing image slots found")
        return
    log.info("Found %d missing image slot(s)", len(slots))

    # (date, hhmm) ごとにまとめる（同じ時刻の複数地点をまとめてGeminiに1回で投げるため）
    by_time = {}
    for date_str, hhmm, name in slots:
        by_time.setdefault((date_str, hhmm), []).append(name)

    today = datetime.now(JST).date()
    api_keys = collect.get_api_keys()
    if not api_keys:
        log.warning("GEMINI_API_KEY is not set; recovered images will be saved but not classified")

    wb = None
    ws = None
    row_index = None
    filled = 0
    recovered_but_invalid = 0
    still_missing = 0
    gave_up = 0
    confidence_updates = {}
    reasoning_updates = {}

    for (date_str, hhmm), names in sorted(by_time.items()):
        target_dt = datetime.strptime(f"{date_str} {hhmm}", "%Y-%m-%d %H%M")
        age_days = (today - target_dt.date()).days
        hour_dir = collect.IMAGE_ARCHIVE_DIR / date_str / hhmm

        recovered_bytes = {}
        for name in names:
            cam_id = cam_id_by_name.get(name)
            if cam_id is None:
                continue
            content, found_dt = fetch_history_image(cam_id, target_dt, station_name=name)
            if content is None:
                if age_days <= GIVE_UP_AFTER_DAYS:
                    log.info("Still unavailable: %s %s %s", date_str, hhmm, name)
                else:
                    gave_up += 1
                continue
            recovered_bytes[name] = content
            offset = int((found_dt - target_dt).total_seconds() / 60)
            log.info("Recovered via history: %s %s %s (offset %+d min)", date_str, hhmm, name, offset)

        if not recovered_bytes:
            still_missing += len(names)
            continue

        # Geminiで再分類（この時刻分だけをまとめて1回のAPI呼び出しにする）
        phenom_by_name, invalid_stations, confidence_by_name, reasoning_by_name = {}, set(), {}, {}
        if api_keys:
            resized_by_name = {}
            for name, raw_bytes in recovered_bytes.items():
                try:
                    resized_by_name[name] = collect.resize_image(raw_bytes)
                except Exception as e:
                    log.warning("Resize failed for recovered image: %s %s %s: %s", date_str, hhmm, name, e)
            if resized_by_name:
                contents, index_to_name = collect._build_contents(list(resized_by_name.items()))
                model = os.environ.get("GEMINI_MODEL", DEFAULT_GEMINI_MODEL)
                parsed = collect._classify_with_model(contents, model, api_keys)
                if parsed is not None:
                    # _parse_classifications()は(result, invalid, low_confidence, confidence,
                    # reasoning)の5値を返す。以前はここが4値で受けており、この再分類パスが
                    # 実行されるたび ValueError で落ちていた(recovered_bytesがあり、かつ
                    # APIキー設定済みで分類に成功した場合に必ず踏む欠陥だったため修正)。
                    phenom_by_name, invalid_stations, _low, confidence_by_name, reasoning_by_name = (
                        collect._parse_classifications(parsed, index_to_name)
                    )

        for name, raw_bytes in recovered_bytes.items():
            if name in invalid_stations:
                recovered_but_invalid += 1
                log.info(
                    "Recovered image judged non-photo/placeholder by Gemini, keeping placeholder: %s %s %s",
                    date_str, hhmm, name,
                )
                continue

            if not dry_run:
                (hour_dir / f"{name}.jpg").write_bytes(raw_bytes)
                (hour_dir / f"{name}_unavailable.jpg").unlink(missing_ok=True)
            log.info("Backfilled image: %s %s %s%s", date_str, hhmm, name, " [dry-run]" if dry_run else "")

            code = phenom_by_name.get(name)
            if code is None:
                continue
            if wb is None:
                wb = openpyxl.load_workbook(DATA_XLSX)
                ws = wb.active
                row_index = _row_index_map(ws)
            row = row_index.get(target_dt)
            if row is None:
                log.warning("No xlsx row found for %s, image saved but phenomenon code not recorded", target_dt)
                continue
            if not dry_run:
                ws.cell(row=row, column=col_by_name[name]).value = "/" if code == 0 else code
            filled += 1
            conf = confidence_by_name.get(name)
            if conf:
                confidence_updates.setdefault(target_dt, {})[name] = conf
                reasoning_updates.setdefault(target_dt, {})[name] = reasoning_by_name.get(name, "")

    if wb is not None and filled and not dry_run:
        wb.save(DATA_XLSX)

    if confidence_updates and not dry_run:
        for dt, conf_map in confidence_updates.items():
            collect.append_confidence_log(conf_map, reasoning_updates.get(dt, {}), dt)

    if gave_up:
        log.info(
            "%d image(s) older than %d days still unavailable; will keep checking quietly",
            gave_up, GIVE_UP_AFTER_DAYS,
        )

    log.info(
        "=== Backfill finished: %d image(s) recovered and classified, %d recovered but judged invalid, "
        "%d still missing ===",
        filled, recovered_but_invalid, still_missing,
    )


if __name__ == "__main__":
    try:
        main()
    except Exception:
        log.exception("backfill_images.py terminated abnormally")
        sys.exit(1)
